"""Comprehensive Automated Unit and Integration Test Suite for Week 8 Day 5 Capstone.
Verifies:
- Documentation files completeness, formatting, and lack of placeholders
- Presentation suite files (PPTX, HTML, Markdown, 1-Pager, Speaker Notes)
- Automated PowerPoint generation with custom styled slides and speaker notes
- Interactive HTML slide deck integrity and embedded ROI calculator logic
- Central metrics consistency across Day 1, Day 2, Day 3, and Day 4
- ROI financial mathematical model calculations
"""
from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

# Add Day 5 root to sys.path
DAY5_DIR = Path(__file__).resolve().parent.parent
if str(DAY5_DIR) not in sys.path:
    sys.path.insert(0, str(DAY5_DIR))

from src.generate_pptx import create_deck
from src.presentation_data import (
    CUSTOMER_PERSONAS,
    EXECUTIVE_METRICS,
    GEOGRAPHIC_SLICES,
    LEAD_BENCHMARK_TABLE,
    PRICE_TIER_SLICES,
    VALUATION_BENCHMARK_TABLE,
    calculate_roi_summary,
)
from src.verify_documentation import (
    EXPECTED_DOC_FILES,
    EXPECTED_PRESENTATION_FILES,
    audit_documentation_suite,
    audit_presentation_suite,
)


class TestDay5Capstone(unittest.TestCase):
    """Test suite verifying Week 8 Day 5 deliverables and artifacts."""

    def setUp(self):
        self.docs_dir = DAY5_DIR / "docs"
        self.pres_dir = DAY5_DIR / "presentation"
        self.assets_dir = self.pres_dir / "assets"

    def test_01_documentation_suite_integrity(self):
        """Verify all 9 documentation files exist, have sufficient depth, and no placeholders."""
        result = audit_documentation_suite(self.docs_dir)
        self.assertTrue(result["passed"], f"Documentation audit failed: {result}")
        self.assertEqual(result["files_checked"], 9)
        self.assertGreater(result["total_words"], 5000)
        self.assertEqual(len(result["missing_files"]), 0)
        self.assertEqual(len(result["placeholder_issues"]), 0)

    def test_02_presentation_suite_files(self):
        """Verify all expected presentation assets exist and have substantial size."""
        result = audit_presentation_suite(self.pres_dir)
        self.assertTrue(result["passed"], f"Presentation audit failed: {result}")
        self.assertEqual(result["files_checked"], 5)
        self.assertEqual(len(result["missing_files"]), 0)

        pptx_file = self.pres_dir / "stakeholder_presentation.pptx"
        html_file = self.pres_dir / "index.html"
        md_file = self.pres_dir / "stakeholder_deck.md"
        summary_file = self.pres_dir / "executive_summary_1pager.md"
        script_file = self.pres_dir / "speaker_notes_script.md"

        self.assertGreater(pptx_file.stat().st_size, 20000)
        self.assertGreater(html_file.stat().st_size, 10000)
        self.assertGreater(md_file.stat().st_size, 5000)
        self.assertGreater(summary_file.stat().st_size, 1000)
        self.assertGreater(script_file.stat().st_size, 5000)

    def test_03_pptx_deck_generation(self):
        """Verify python-pptx generation generates a valid deck with >= 14 slides and speaker notes."""
        with tempfile.TemporaryDirectory() as tmp_dir:
            tmp_pptx = Path(tmp_dir) / "test_deck.pptx"
            gen_path = create_deck(tmp_pptx)
            self.assertTrue(gen_path.exists())
            self.assertGreater(gen_path.stat().st_size, 15000)

            # Inspect presentation structure
            from pptx import Presentation
            prs = Presentation(str(gen_path))
            self.assertGreaterEqual(len(prs.slides), 14)

            # Verify speaker notes are populated on slides
            notes_count = 0
            for slide in prs.slides:
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
                    notes_count += 1
            self.assertGreaterEqual(notes_count, 12, "At least 12 slides must contain speaker notes.")

    def test_04_html_presentation_structure(self):
        """Verify interactive HTML presentation has slides, keyboard controls, and interactive elements."""
        html_path = self.pres_dir / "index.html"
        content = html_path.read_text(encoding="utf-8")

        self.assertIn("AI Property Valuation", content)
        self.assertIn("roi-calc-container", content)
        self.assertIn("speakerNotesPanel", content)
        self.assertIn("btnFullscreen", content)
        self.assertIn("updateRoiCalc", content)
        self.assertIn("Outfit", content)
        self.assertIn("Inter", content)
        # Check all 14 slide containers exist
        for i in range(1, 15):
            self.assertIn(f'id="slide-{i}"', content)

    def test_05_central_metrics_consistency(self):
        """Verify centralized metrics agree with Day 2, Day 3, and Day 4 specifications."""
        self.assertEqual(EXECUTIVE_METRICS["total_listings_trained"], 10480)
        self.assertEqual(EXECUTIVE_METRICS["total_leads_analyzed"], 3496)
        self.assertAlmostEqual(EXECUTIVE_METRICS["valuation_r2"], 0.9860, places=3)
        self.assertAlmostEqual(EXECUTIVE_METRICS["valuation_mape"], 0.0729, places=3)
        self.assertAlmostEqual(EXECUTIVE_METRICS["lead_roc_auc"], 0.8201, places=3)
        self.assertAlmostEqual(EXECUTIVE_METRICS["lead_f1_score"], 0.6136, places=3)
        self.assertAlmostEqual(EXECUTIVE_METRICS["precision_top_20"], 0.7143, places=3)
        self.assertAlmostEqual(EXECUTIVE_METRICS["precision_top_20_lift"], 2.50, places=2)
        self.assertEqual(EXECUTIVE_METRICS["voice_agent_sla_mins"], 15)

    def test_06_roi_financial_calculation_math(self):
        """Verify ROI formulas produce mathematically sound economics."""
        roi_default = calculate_roi_summary(monthly_inbound_leads=1000)
        self.assertEqual(roi_default["monthly_inbound_leads"], 1000)
        self.assertEqual(roi_default["monthly_leads_called"], 200)
        self.assertAlmostEqual(roi_default["baseline_monthly_deals"], 57.1, places=1)
        self.assertAlmostEqual(roi_default["ai_monthly_deals"], 142.9, places=1)
        self.assertAlmostEqual(roi_default["incremental_monthly_deals"], 85.7, places=1)
        self.assertGreater(roi_default["annual_incremental_commission_pkr"], 30000000.0)
        self.assertGreater(roi_default["roi_multiple"], 10.0)

    def test_07_visual_assets_presence(self):
        """Verify that essential figures from earlier days are available in Day 5 presentation assets."""
        self.assertTrue(self.assets_dir.exists())
        required_figures = [
            "actual_vs_predicted.png",
            "error_slices_breakdown.png",
            "confusion_matrix_top_models.png",
            "roc_pr_curves_comparison.png",
            "shap_global_importance.png",
            "customer_personas_clusters.png",
        ]
        for fig in required_figures:
            fig_path = self.assets_dir / fig
            self.assertTrue(fig_path.exists(), f"Figure missing from Day 5 assets: {fig}")
            self.assertGreater(fig_path.stat().st_size, 5000)

    def test_08_customer_personas_and_slices_completeness(self):
        """Verify personas and geographic slice metadata."""
        self.assertEqual(len(CUSTOMER_PERSONAS), 4)
        self.assertEqual(len(GEOGRAPHIC_SLICES), 4)
        self.assertEqual(len(PRICE_TIER_SLICES), 4)

        # Check persona names
        persona_names = [p["name"] for p in CUSTOMER_PERSONAS]
        self.assertIn("Overseas Capital Investor", persona_names)
        self.assertIn("Urgent Family Homebuyer", persona_names)

        # Check city names
        cities = [g["city"] for g in GEOGRAPHIC_SLICES]
        self.assertIn("Islamabad", cities)
        self.assertIn("Karachi", cities)
        self.assertIn("Lahore", cities)
        self.assertIn("Rawalpindi", cities)


if __name__ == "__main__":
    unittest.main()
